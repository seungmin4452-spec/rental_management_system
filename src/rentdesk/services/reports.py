"""ReportService: 5개 표준 엑셀 보고서 (PRD P0-10, 12장, S-09).

1. 장비 목록   (카테고리, 운영 상태)
2. 예약 목록   (시작/끝 일자, 주문 상태)
3. 오늘 출고·반납 (기준일)
4. 연체 목록   (기준 시각)
5. 정산 목록   (기간)

보고서 데이터(`ReportTable`) 생성과 xlsx 쓰기를 분리해 내용은 순수 데이터로 테스트한다.
시트명·제목은 한국어, 일시는 서울 시각 `yyyy-mm-dd hh:mm`, 금액은 `#,##0"원"` 서식.
금융/세무 공식 증빙이 아니라는 안내를 모든 보고서에 넣는다.
"""

from __future__ import annotations

import os
import sqlite3
import tempfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from enum import StrEnum
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rentdesk.domain.availability import ASSET_STATE_LABELS
from rentdesk.domain.display import format_duration, seoul_day_range
from rentdesk.domain.errors import DomainError, ValidationError
from rentdesk.domain.models import (
    ORDER_STATUS_LABELS,
    LifecycleStatus,
    OrderStatus,
)
from rentdesk.domain.timeutil import SEOUL, format_seoul, to_utc
from rentdesk.persistence import repositories as repo
from rentdesk.services.assets import AssetService
from rentdesk.services.availability import AvailabilityService
from rentdesk.services.common import Clock, current_time, default_clock
from rentdesk.services.dashboard import overdue_rows, schedule_rows
from rentdesk.services.reservations import ReservationService
from rentdesk.services.settings import BusinessInfo, read_business_info
from rentdesk.services.settlements import MONEY_ALLOWED_STATUSES, load_summary

DISCLAIMER = "본 자료는 내부 관리용이며 금융·세무 공식 증빙(영수증·세금계산서)이 아닙니다."
WON_FORMAT = '#,##0"원"'
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
LIFECYCLE_LABELS = {LifecycleStatus.ACTIVE: "활성", LifecycleStatus.RETIRED: "은퇴"}


class ColumnType(StrEnum):
    TEXT = "TEXT"
    INT = "INT"
    WON = "WON"
    DATETIME = "DATETIME"


@dataclass(frozen=True)
class Column:
    header: str
    type: ColumnType = ColumnType.TEXT
    width: int = 14


@dataclass(frozen=True)
class ReportTable:
    title: str
    sheet_name: str
    file_stem: str
    filters: tuple[str, ...]
    columns: tuple[Column, ...]
    rows: tuple[tuple[object, ...], ...]
    totals: tuple[tuple[str, int], ...] = field(default=())  # (라벨, 금액) 하단 합계


class ReportExistsError(DomainError):
    """같은 이름의 파일이 이미 있음. 화면은 덮어쓰기 확인 후 overwrite=True로 다시 호출한다."""


def _validate_range(start: date, end: date) -> None:
    if not isinstance(start, date) or not isinstance(end, date):
        raise ValidationError("period", "조회 기간을 입력해주세요.")
    if start > end:
        raise ValidationError("period", "조회 시작일은 종료일보다 늦을 수 없습니다.")


def _period_utc(start: date, end: date) -> tuple[datetime, datetime]:
    """서울 날짜 [start, end] (종료일 포함) → UTC `[start 00:00, end+1 00:00)`."""
    return seoul_day_range(start)[0], seoul_day_range(end)[1]


def _stamp(d: date) -> str:
    return d.strftime("%Y%m%d")


class ReportService:
    def __init__(self, conn: sqlite3.Connection, clock: Clock = default_clock) -> None:
        self._conn = conn
        self._clock = clock

    # ------------------------------------------------------------------ builders

    def assets_report(
        self, *, category_id: int | None = None, lifecycle: LifecycleStatus | None = None
    ) -> ReportTable:
        assets = AssetService(self._conn, self._clock).search_assets(
            include_retired=lifecycle != LifecycleStatus.ACTIVE, category_id=category_id
        )
        if lifecycle is not None:
            assets = [a for a in assets if a.lifecycle_status == lifecycle]
        states = AvailabilityService(self._conn, self._clock).asset_states(assets)
        filters = [
            "카테고리: " + (
                repo.get_category(self._conn, category_id).name if category_id else "전체"
            ),
            "운영 상태: " + (LIFECYCLE_LABELS[lifecycle] if lifecycle else "전체"),
        ]
        today = current_time(self._clock).astimezone(SEOUL).date()
        return ReportTable(
            title="장비 목록",
            sheet_name="장비 목록",
            file_stem=f"장비목록_{_stamp(today)}",
            filters=tuple(filters),
            columns=(
                Column("장비번호", width=14), Column("장비명", width=26),
                Column("카테고리", width=12), Column("브랜드", width=12),
                Column("모델", width=16), Column("일 대여료", ColumnType.WON),
                Column("보증금", ColumnType.WON), Column("현재 상태", width=10),
                Column("운영 상태", width=10), Column("보관 위치", width=12),
            ),
            rows=tuple(
                (
                    a.asset_code, a.name, a.category_name, a.brand or "", a.model or "",
                    a.daily_rate_won, a.deposit_won, ASSET_STATE_LABELS[states[a.id]],
                    LIFECYCLE_LABELS[a.lifecycle_status], a.location or "",
                )
                for a in assets
            ),
        )

    def reservations_report(
        self, start: date, end: date, *, statuses: Sequence[OrderStatus] | None = None
    ) -> ReportTable:
        _validate_range(start, end)
        period_start, period_end = _period_utc(start, end)
        rows = ReservationService(self._conn, self._clock).list_order_rows(
            statuses=statuses, period_start=period_start, period_end=period_end
        )
        rows.sort(key=lambda r: (r.order.start_at, r.order.order_no))
        status_desc = ", ".join(ORDER_STATUS_LABELS[s] for s in statuses) if statuses else "전체"
        return ReportTable(
            title="예약 목록",
            sheet_name="예약 목록",
            file_stem=f"예약목록_{_stamp(start)}-{_stamp(end)}",
            filters=(f"기간: {start} ~ {end} (예약 기간이 겹치는 예약)", f"상태: {status_desc}"),
            columns=(
                Column("예약번호", width=17), Column("고객", width=18),
                Column("장비 수", ColumnType.INT, 8), Column("예약 시작", ColumnType.DATETIME, 17),
                Column("반납 예정", ColumnType.DATETIME, 17), Column("상태", width=10),
                Column("출고/반납", width=14), Column("예상 대여료", ColumnType.WON),
                Column("예상 보증금", ColumnType.WON),
            ),
            rows=tuple(
                (
                    r.order.order_no, r.customer_name, r.item_count, r.order.start_at,
                    r.order.due_at, ORDER_STATUS_LABELS[r.order.status], r.progress_label,
                    r.expected_total_rental_fee_won, r.expected_deposit_won,
                )
                for r in rows
            ),
        )

    def daily_schedule_report(self, day: date) -> ReportTable:
        if not isinstance(day, date):
            raise ValidationError("day", "기준일을 입력해주세요.")
        checkouts, returns = schedule_rows(self._conn, day)
        rows = [("출고", r) for r in checkouts] + [("반납", r) for r in returns]
        return ReportTable(
            title="오늘 출고·반납",
            sheet_name="출고·반납",
            file_stem=f"출고반납_{_stamp(day)}",
            filters=(f"기준일: {day}",),
            columns=(
                Column("구분", width=6), Column("예약번호", width=17),
                Column("장비번호", width=14), Column("장비명", width=24),
                Column("고객", width=18), Column("예정 시각", ColumnType.DATETIME, 17),
                Column("처리 상태", width=10),
            ),
            rows=tuple(
                (
                    kind, r.order_no, r.asset_code, r.asset_name, r.customer_name,
                    r.scheduled_at, r.state_label,
                )
                for kind, r in rows
            ),
        )

    def overdue_report(self, at: datetime | None = None) -> ReportTable:
        at = to_utc(at) if at is not None else current_time(self._clock)
        rows = overdue_rows(self._conn, at)
        return ReportTable(
            title="연체 목록",
            sheet_name="연체 목록",
            file_stem=f"연체목록_{at.astimezone(SEOUL).strftime('%Y%m%d_%H%M')}",
            filters=(
                f"기준 시각: {format_seoul(at)}",
                "예상 연체료는 기준 시각에 반납한다고 가정한 예측값이며 아직 청구되지 않았습니다.",
            ),
            columns=(
                Column("고객", width=18), Column("예약번호", width=17),
                Column("장비번호", width=14), Column("장비명", width=24),
                Column("약정 반납", ColumnType.DATETIME, 17), Column("지연 시간", width=14),
                Column("예상 연체료", ColumnType.WON),
            ),
            rows=tuple(
                (
                    r.customer_name, r.order_no, r.asset_code, r.asset_name, r.due_at,
                    format_duration(r.overdue_seconds), r.expected_late_fee_won,
                )
                for r in rows
            ),
            totals=(("예상 연체료 합계", sum(r.expected_late_fee_won for r in rows)),),
        )

    def settlement_report(
        self, start: date, end: date, *, unsettled_only: bool = False
    ) -> ReportTable:
        _validate_range(start, end)
        period_start, period_end = _period_utc(start, end)
        orders = ReservationService(self._conn, self._clock).list_order_rows(
            statuses=sorted(MONEY_ALLOWED_STATUSES), period_start=period_start,
            period_end=period_end,
        )
        orders.sort(key=lambda r: (r.order.start_at, r.order.order_no))
        data = []
        for r in orders:
            s = load_summary(self._conn, r.order.id)
            if unsettled_only and s.is_settled:
                continue
            data.append((r, s))
        return ReportTable(
            title="정산 목록",
            sheet_name="정산 목록",
            file_stem=f"정산목록_{_stamp(start)}-{_stamp(end)}",
            filters=(
                f"기간: {start} ~ {end} (출고 이후 예약 중 예약 기간이 겹치는 예약)",
                "범위: " + ("미정산만" if unsettled_only else "전체"),
            ),
            columns=(
                Column("예약번호", width=17), Column("고객", width=18),
                Column("상태", width=10), Column("출고 시각", ColumnType.DATETIME, 17),
                Column("청구 합계", ColumnType.WON), Column("대여료 순수금", ColumnType.WON),
                Column("보증금 대금충당", ColumnType.WON), Column("미수금", ColumnType.WON),
                Column("보증금 보유액", ColumnType.WON), Column("정산", width=8),
            ),
            rows=tuple(
                (
                    r.order.order_no, r.customer_name, ORDER_STATUS_LABELS[r.order.status],
                    r.order.actual_checkout_at, s.charge_total_won, s.net_rental_receipts_won,
                    s.net_deposit_applied_won, s.amount_due_won, s.deposit_held_won,
                    "완료" if s.is_settled else "미정산",
                )
                for r, s in data
            ),
            totals=(
                ("청구 합계", sum(s.charge_total_won for _, s in data)),
                ("대여료 순수금 합계", sum(s.net_rental_receipts_won for _, s in data)),
                ("미수금 합계", sum(s.amount_due_won for _, s in data)),
                ("보증금 보유액 합계", sum(s.deposit_held_won for _, s in data)),
            ),
        )

    # ------------------------------------------------------------------ writer

    def export(self, table: ReportTable, path: Path | str, *, overwrite: bool = False) -> Path:
        """xlsx로 저장. 같은 폴더 임시 파일에 쓴 뒤 교체해 중간 실패 시 기존 파일을 보존한다."""
        target = Path(path)
        if target.suffix.lower() != ".xlsx":
            target = target.with_name(target.name + ".xlsx")
        if target.exists() and not overwrite:
            raise ReportExistsError(f"같은 이름의 파일이 이미 있습니다: {target.name}")
        if not target.parent.is_dir():
            raise DomainError("저장할 폴더를 찾을 수 없습니다. 다른 위치를 선택해주세요.")
        workbook = build_workbook(
            table, read_business_info(self._conn), current_time(self._clock)
        )
        tmp_name = None
        try:
            fd, tmp_name = tempfile.mkstemp(
                prefix=".rentdesk-", suffix=".xlsx", dir=str(target.parent)
            )
            os.close(fd)
            workbook.save(tmp_name)
            os.replace(tmp_name, target)
            tmp_name = None
        except PermissionError as exc:
            raise DomainError(
                "파일을 저장할 수 없습니다. 같은 파일이 엑셀 등 다른 프로그램에서 열려 있거나 "
                "폴더 쓰기 권한이 없습니다."
            ) from exc
        except OSError as exc:
            raise DomainError(
                "파일을 저장하지 못했습니다. 저장 위치와 남은 용량을 확인해주세요."
            ) from exc
        finally:
            if tmp_name is not None and os.path.exists(tmp_name):
                os.remove(tmp_name)
        return target


def _excel_value(value: object, col: Column) -> object:
    if col.type == ColumnType.DATETIME:
        if value is None:
            return None
        assert isinstance(value, datetime)
        return value.astimezone(SEOUL).replace(tzinfo=None)  # 엑셀은 시간대 없는 서울 시각
    return value


def build_workbook(table: ReportTable, business: BusinessInfo, generated_at: datetime) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = table.sheet_name[:31]
    ncols = len(table.columns)

    biz = " / ".join(x for x in (business.name, business.phone) if x)
    meta = [f"작성: {format_seoul(generated_at)}"] + ([f"사업자: {biz}"] if biz else [])
    ws.append([table.title])
    ws.cell(row=1, column=1).font = Font(size=14, bold=True)
    ws.append([" · ".join(meta)])
    for f in table.filters:
        ws.append([f])
    ws.append([DISCLAIMER])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="666666")
    ws.append([])

    header_row = ws.max_row + 1
    ws.append([c.header for c in table.columns])
    fill = PatternFill("solid", fgColor="DDE6F0")
    for idx in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=idx)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for row in table.rows:
        ws.append([_excel_value(v, c) for v, c in zip(row, table.columns, strict=True)])
        r = ws.max_row
        for idx, col in enumerate(table.columns, start=1):
            if col.type == ColumnType.WON:
                ws.cell(row=r, column=idx).number_format = WON_FORMAT
            elif col.type == ColumnType.DATETIME:
                ws.cell(row=r, column=idx).number_format = DATETIME_FORMAT

    if not table.rows:
        ws.append(["조건에 해당하는 자료가 없습니다."])
    if table.totals:
        ws.append([])
        for label, amount in table.totals:
            ws.append([label, amount])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=2).number_format = WON_FORMAT

    for idx, col in enumerate(table.columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = col.width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if table.rows:
        last = header_row + len(table.rows)
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{last}"
    return wb
