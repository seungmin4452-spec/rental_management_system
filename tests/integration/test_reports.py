"""P0-10 / PRD 12장: 5개 표준 엑셀 보고서 내용·서식·파일 저장 안전성."""

import os
import stat
import sys
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from rentdesk.domain.errors import DomainError, ValidationError
from rentdesk.domain.models import LifecycleStatus, MoneyKind, OrderStatus
from rentdesk.services.reports import (
    DATETIME_FORMAT,
    DISCLAIMER,
    WON_FORMAT,
    ReportExistsError,
)
from rentdesk.services.settlements import MoneyInput
from tests.support import kst


@pytest.fixture
def data(svc):
    cam = svc.add_asset("CAM-1", rate=30_000, deposit=100_000)
    len_ = svc.add_asset("LEN-1", rate=10_000, deposit=30_000, category="렌즈", name="표준 렌즈")
    old = svc.add_asset("OLD-1", category="렌즈", name="구형 렌즈")
    svc.assets.retire_asset(old.id)
    kim = svc.add_customer("김하늘")
    lee = svc.add_customer("이바다")
    o1 = svc.book(kim.id, [cam.id, len_.id], kst(10, 19, 10), kst(10, 20, 10))
    o2 = svc.book(lee.id, [cam.id], kst(10, 23, 10), kst(10, 24, 10))
    svc.checkout(o1.id, kst(10, 19, 10), deposit_in_won=130_000)
    svc.return_asset(o1.id, len_.id, kst(10, 20, 9))
    svc.clock.now = kst(10, 21, 12)
    return dict(cam=cam, len=len_, old=old, o1=o1, o2=o2)


def test_assets_report_filters(svc, data):
    all_rows = svc.reports.assets_report()
    assert [r[0] for r in all_rows.rows] == ["CAM-1", "LEN-1", "OLD-1"]
    assert all_rows.rows[0][5:9] == (30_000, 100_000, "연체", "운영 중")

    lens = next(c for c in svc.assets.list_categories() if c.name == "렌즈")
    lens_active = svc.reports.assets_report(category_id=lens.id, lifecycle=LifecycleStatus.ACTIVE)
    assert [r[0] for r in lens_active.rows] == ["LEN-1"]
    retired = svc.reports.assets_report(lifecycle=LifecycleStatus.RETIRED)
    assert [r[0] for r in retired.rows] == ["OLD-1"]
    assert "운영 상태: 사용 중지" in retired.filters


def test_reservations_report_period_and_status(svc, data):
    t = svc.reports.reservations_report(date(2026, 10, 1), date(2026, 10, 31))
    assert [r[0] for r in t.rows] == [data["o1"].order_no, data["o2"].order_no]
    assert t.rows[0][1:3] == ("김하늘", 2)
    assert t.rows[0][3] == kst(10, 19, 10)

    only_confirmed = svc.reports.reservations_report(
        date(2026, 10, 1), date(2026, 10, 31), statuses=[OrderStatus.CONFIRMED]
    )
    assert [r[0] for r in only_confirmed.rows] == [data["o2"].order_no]
    # 종료일 포함: 10/23 하루만 조회해도 10/23 시작 예약이 포함된다
    one_day = svc.reports.reservations_report(date(2026, 10, 23), date(2026, 10, 23))
    assert [r[0] for r in one_day.rows] == [data["o2"].order_no]

    with pytest.raises(ValidationError):
        svc.reports.reservations_report(date(2026, 10, 31), date(2026, 10, 1))


def test_daily_schedule_report(svc, data):
    t = svc.reports.daily_schedule_report(date(2026, 10, 20))
    assert [(r[0], r[2], r[6]) for r in t.rows] == [
        ("반납", "CAM-1", "대기"),
        ("반납", "LEN-1", "처리 완료"),
    ]
    t19 = svc.reports.daily_schedule_report(date(2026, 10, 19))
    assert {(r[0], r[2], r[6]) for r in t19.rows} == {("출고", "CAM-1", "처리 완료"),
                                                      ("출고", "LEN-1", "처리 완료")}


def test_overdue_report_at_reference_time(svc, data):
    t = svc.reports.overdue_report(kst(10, 21, 12))
    assert len(t.rows) == 1
    customer, order_no, code, _name, due, delay, fee = t.rows[0]
    assert (customer, order_no, code) == ("김하늘", data["o1"].order_no, "CAM-1")
    assert due == kst(10, 20, 10)
    assert delay == "1일 2시간"
    assert fee == 60_000  # 26h - 60분 유예 → 2일 × 30,000원
    assert t.totals == (("예상 연체료 합계", 60_000),)
    # 기준 시각 이전에는 연체 아님
    assert svc.reports.overdue_report(kst(10, 20, 10)).rows == ()


def test_settlement_report(svc, data):
    t = svc.reports.settlement_report(date(2026, 10, 1), date(2026, 10, 31))
    assert len(t.rows) == 1  # 출고 전 예약(o2)은 제외
    row = t.rows[0]
    assert row[0] == data["o1"].order_no
    assert row[4:11] == (40_000, 0, 0, 0, 40_000, 130_000, "미정산")
    assert ("미수금 합계", 40_000) in t.totals

    svc.settlements.record(data["o1"].id, MoneyInput(kind=MoneyKind.DEPOSIT_APPLIED, amount_won=40_000))
    svc.settlements.record(data["o1"].id, MoneyInput(kind=MoneyKind.DEPOSIT_OUT, amount_won=90_000))
    assert svc.reports.settlement_report(
        date(2026, 10, 1), date(2026, 10, 31), unsettled_only=True
    ).rows == ()


def test_export_xlsx_korean_titles_and_formats(svc, data, tmp_path):
    svc.settings.update_business_info(name="예제 렌탈", phone="02-000-0000")
    folder = tmp_path / "보고서 폴더 📁"
    folder.mkdir()
    table = svc.reports.overdue_report(kst(10, 21, 12))
    saved = svc.reports.export(table, folder / "연체 목록")

    assert saved.name == "연체 목록.xlsx"
    wb = load_workbook(saved)
    ws = wb["연체 목록"]
    texts = [c.value for c in ws["A"] if c.value]
    assert texts[0] == "연체 목록"
    assert any("예제 렌탈" in str(t) for t in texts)
    assert DISCLAIMER in texts
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "고객")
    data_row = header_row + 1
    assert ws.cell(data_row, 5).value == datetime(2026, 10, 20, 10, 0)  # 서울 시각(시간대 없음)
    assert ws.cell(data_row, 5).number_format == DATETIME_FORMAT
    assert ws.cell(data_row, 7).value == 60_000
    assert ws.cell(data_row, 7).number_format == WON_FORMAT
    assert not [p for p in folder.iterdir() if p.name.startswith(".rentdesk-")]  # 임시 파일 정리


def test_export_all_five_reports(svc, data, tmp_path):
    r = svc.reports
    tables = [
        r.assets_report(),
        r.reservations_report(date(2026, 10, 1), date(2026, 10, 31)),
        r.daily_schedule_report(date(2026, 10, 20)),
        r.overdue_report(),
        r.settlement_report(date(2026, 10, 1), date(2026, 10, 31)),
    ]
    names = set()
    for t in tables:
        path = r.export(t, tmp_path / f"{t.file_stem}.xlsx")
        names.add(load_workbook(path).sheetnames[0])
    assert names == {"장비 목록", "예약 목록", "출고·반납", "연체 목록", "정산 목록"}


def test_export_empty_report_has_notice(svc, tmp_path):
    path = svc.reports.export(svc.reports.overdue_report(), tmp_path / "empty.xlsx")
    ws = load_workbook(path).active
    assert "조건에 해당하는 자료가 없습니다." in [c.value for c in ws["A"]]


def test_export_requires_overwrite_confirmation(svc, data, tmp_path):
    target = tmp_path / "장비.xlsx"
    target.write_bytes(b"old")
    with pytest.raises(ReportExistsError):
        svc.reports.export(svc.reports.assets_report(), target)
    assert target.read_bytes() == b"old"
    svc.reports.export(svc.reports.assets_report(), target, overwrite=True)
    assert load_workbook(target).sheetnames == ["장비 목록"]


def test_export_missing_folder_message(svc, tmp_path):
    with pytest.raises(DomainError, match="폴더를 찾을 수 없습니다"):
        svc.reports.export(svc.reports.assets_report(), tmp_path / "없음" / "a.xlsx")


@pytest.mark.skipif(sys.platform == "win32", reason="POSIX 권한 비트로 쓰기 거부를 재현")
def test_export_permission_error_is_friendly_and_keeps_existing(svc, tmp_path):
    folder = tmp_path / "ro"
    folder.mkdir()
    existing = folder / "a.xlsx"
    existing.write_bytes(b"keep")
    os.chmod(folder, stat.S_IRUSR | stat.S_IXUSR)
    try:
        with pytest.raises(DomainError, match="저장할 수 없습니다"):
            svc.reports.export(svc.reports.assets_report(), existing, overwrite=True)
    finally:
        os.chmod(folder, stat.S_IRWXU)
    assert existing.read_bytes() == b"keep"


def test_settlement_report_includes_canceled_with_money_only(svc, data):
    """U5: 돈이 오간 취소 예약은 정산 목록에 포함, 돈이 없던 취소 예약은 제외."""
    svc.settlements.record(data["o2"].id, MoneyInput(kind=MoneyKind.RENTAL_RECEIPT, amount_won=10_000))
    svc.reservations.cancel(data["o2"].id, cancel_fee_won=3_000)
    t = svc.reports.settlement_report(date(2026, 10, 1), date(2026, 10, 31))
    assert data["o2"].order_no in [r[0] for r in t.rows]
