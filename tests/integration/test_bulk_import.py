"""엑셀 일괄 가져오기: 장비·고객 (USABILITY_PLAN U10)."""

from __future__ import annotations

import time

import pytest
from openpyxl import Workbook, load_workbook

from rentdesk.domain.errors import DomainError, InvalidStateError, ValidationError
from rentdesk.domain.models import CustomerType
from rentdesk.services.backup import inspect_backup
from rentdesk.services.bulk_import import (
    ASSET_COLUMNS,
    BACKUP_PREFIX,
    CUSTOMER_COLUMNS,
    ImportKind,
    ImportService,
    template_name,
)
from rentdesk.services.customers import CustomerInput

ASSET_HEADERS = ["관리번호", "장비명", "카테고리", "브랜드", "모델명", "시리얼 번호", "일 대여료", "보증금"]
CUSTOMER_HEADERS = ["이름", "구분", "전화번호", "이메일", "메모"]


def write_xlsx(path, headers, rows, *, sheet="장비", extra_sheets=()):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    for name in extra_sheets:
        wb.create_sheet(name)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture
def imports(conn, clock):
    return ImportService(conn, clock)


@pytest.fixture
def folder(tmp_path):
    f = tmp_path / "가져오기 폴더"
    f.mkdir()
    return f


# --------------------------------------------------------------------------- 양식


@pytest.mark.parametrize("kind", list(ImportKind))
def test_template_has_header_example_and_guide(imports, folder, kind):
    path = imports.save_template(kind, folder / template_name(kind))
    wb = load_workbook(path)
    assert wb.sheetnames == [("장비" if kind == ImportKind.ASSETS else "고객"), "설명"]
    ws = wb.worksheets[0]
    columns = ASSET_COLUMNS if kind == ImportKind.ASSETS else CUSTOMER_COLUMNS
    assert [c.value for c in ws[1]] == [c.template_header for c in columns]
    assert [c.value for c in ws[2]] == [c.example for c in columns]
    guide_text = " ".join(str(c.value) for row in wb["설명"].iter_rows() for c in row if c.value)
    assert "필수" in guide_text and "백업" in guide_text


@pytest.mark.parametrize("kind", list(ImportKind))
def test_untouched_template_imports_nothing(imports, folder, kind):
    """양식을 받아 예시만 둔 채 고르면 예시 행은 가져오지 않는다."""
    path = imports.save_template(kind, folder / template_name(kind))
    preview = imports.preview(kind, path)
    assert preview.rows == ()
    assert any("예시" in n for n in preview.notices)
    assert any("가져올 행이 없습니다" in n for n in preview.notices)


def test_filled_template_roundtrip(imports, conn, folder):
    path = imports.save_template(ImportKind.ASSETS, folder / template_name(ImportKind.ASSETS))
    wb = load_workbook(path)
    ws = wb["장비"]
    ws.append(["LENS-01", "망원 렌즈", "렌즈", None, None, None, 20000, None, None, None])
    wb.save(path)
    preview = imports.preview(ImportKind.ASSETS, path)
    assert [r.row_no for r in preview.rows] == [3]  # 2행 예시는 건너뜀
    result = imports.run_import(preview, create_categories=True, backup_dir=None)
    assert result.created == 1


# --------------------------------------------------------------------------- 장비


def test_assets_normal_import(imports, svc, folder):
    svc.add_asset("CAM-1", category="카메라")
    path = write_xlsx(folder / "장비 목록.xlsx", ASSET_HEADERS, [
        ("CAM-2", "카메라 B", "카메라", "소니", "A7M4", 12345, "50,000원", 300000),
        (None, "삼각대", "카메라", None, None, None, 10000, None),
    ])
    preview = imports.preview(ImportKind.ASSETS, path)
    assert [r.is_ok(True) for r in preview.rows] == [True, True]
    assert preview.new_categories == ()
    result = imports.run_import(preview, create_categories=False, backup_dir=None)
    assert (result.created, result.skipped) == (2, 0)
    assets = {a.asset_code: a for a in svc.assets.search_assets()}
    b = assets["CAM-2"]
    assert (b.daily_rate_won, b.deposit_won, b.serial_no, b.brand) == (50000, 300000, "12345", "소니")
    auto = [a for code, a in assets.items() if code.startswith("AST-")]
    assert len(auto) == 1 and auto[0].deposit_won == 0  # 빈 관리번호 자동 부여, 빈 보증금 0원


def test_assets_duplicates_ignore_case(imports, svc, folder):
    svc.add_asset("CAM-1")
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [
        ("cam-1", "중복(DB)", "카메라", None, None, None, 1000, 0),
        ("LENS-1", "렌즈", "카메라", None, None, None, 1000, 0),
        ("lens-1", "중복(파일)", "카메라", None, None, None, 1000, 0),
    ])
    preview = imports.preview(ImportKind.ASSETS, path)
    r_db, r_ok, r_file = preview.rows
    assert "이미 등록된 관리번호" in r_db.errors[0]
    assert r_ok.is_ok(False)
    assert "3행과 겹칩니다" in r_file.errors[0]  # 엑셀 행 번호(머리글이 1행)
    result = imports.run_import(preview, create_categories=False, backup_dir=None)
    assert (result.created, result.skipped) == (1, 2)


def test_assets_blank_non_numeric_negative(imports, folder):
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [
        ("A-1", None, "카메라", None, None, None, 1000, 0),          # 빈 장비명
        ("A-2", "장비", None, None, None, None, None, 0),            # 빈 카테고리·빈 대여료
        ("A-3", "장비", "카메라", None, None, None, "오만원", 0),     # 숫자 아님
        ("A-4", "장비", "카메라", None, None, None, -5000, -1),       # 음수
        ("A-5", "장비", "카메라", None, None, None, 1000.5, 0),       # 소수
    ])
    rows = imports.preview(ImportKind.ASSETS, path).rows
    assert any("장비명" in e for e in rows[0].errors)
    assert len(rows[1].errors) == 2  # 카테고리 + 일 대여료를 한꺼번에 알려준다
    assert "숫자로" in rows[2].errors[0]
    assert len(rows[3].errors) == 2 and all("0 이상" in e for e in rows[3].errors)
    assert "소수" in rows[4].errors[0]


def test_new_category_choice(imports, svc, folder):
    svc.assets.create_category("카메라")
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [
        ("D-1", "드론 A", "드론", None, None, None, 1000, 0),
        ("D-2", "드론 B", "드론", None, None, None, 1000, 0),
        ("C-1", "카메라", "카메라", None, None, None, 1000, 0),
    ])
    preview = imports.preview(ImportKind.ASSETS, path)
    assert preview.new_categories == ("드론",)
    assert [r.is_ok(False) for r in preview.rows] == [False, False, True]
    assert "새 카테고리" in preview.rows[0].problems(False)[0]

    only_existing = imports.run_import(preview, create_categories=False, backup_dir=None)
    assert only_existing.created == 1 and only_existing.created_categories == ()

    preview = imports.preview(ImportKind.ASSETS, path)  # C-1은 이제 중복
    result = imports.run_import(preview, create_categories=True, backup_dir=None)
    assert result.created == 2 and result.created_categories == ("드론",)
    names = [c.name for c in svc.assets.list_categories()]
    assert names.count("드론") == 1


def test_archived_category_is_error(imports, svc, folder):
    cat = svc.assets.create_category("옛 장비")
    svc.assets.set_category_archived(cat.id, True)
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [("X-1", "옛 것", "옛 장비", None, None, None, 1, 0)])
    row = imports.preview(ImportKind.ASSETS, path).rows[0]
    assert "숨김 처리된 카테고리" in row.errors[0]


def test_auto_code_skips_codes_in_file(imports, svc, folder):
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [
        (None, "자동", "카메라", None, None, None, 1, 0),
        ("ast-000001", "직접", "카메라", None, None, None, 1, 0),
    ])
    imports.run_import(imports.preview(ImportKind.ASSETS, path), create_categories=True, backup_dir=None)
    codes = sorted(a.asset_code for a in svc.assets.search_assets())
    assert codes == ["AST-000002", "ast-000001"]


def test_header_aliases_order_and_unknown_columns(imports, folder):
    path = write_xlsx(folder / "a.xlsx", ["비고", "일 대여료(필수)", "품명", "분류", "장비번호", "구입일"], [
        ("메모", 7000, "조명", "조명", "L-1", "2024-01-01"),
    ])
    preview = imports.preview(ImportKind.ASSETS, path)
    row = preview.rows[0]
    assert row.is_ok(True)
    assert (row.asset.asset_code, row.asset.name, row.asset.note, row.asset.daily_rate_won) == ("L-1", "조명", "메모", 7000)
    assert any("구입일" in n for n in preview.notices)


def test_wrong_template_or_bad_file(imports, folder):
    customers = write_xlsx(folder / "고객.xlsx", CUSTOMER_HEADERS, [("홍길동",)], sheet="고객")
    with pytest.raises(DomainError, match="장비 머리글"):
        imports.preview(ImportKind.ASSETS, customers)
    broken = folder / "깨진 파일.xlsx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(DomainError, match="읽지 못했습니다"):
        imports.preview(ImportKind.ASSETS, broken)
    old = folder / "old.xls"
    old.write_bytes(b"x")
    with pytest.raises(DomainError, match=r"\.xls"):
        imports.preview(ImportKind.ASSETS, old)
    with pytest.raises(DomainError, match="찾을 수 없습니다"):
        imports.preview(ImportKind.ASSETS, folder / "없음.xlsx")


def test_korean_sheet_and_path_with_other_sheets(imports, tmp_path):
    path = write_xlsx(tmp_path / "한글 폴더 (1)" / "우리 가게 장비.xlsx", ASSET_HEADERS,
                      [("K-1", "한글 장비", "한글 분류", None, None, None, 1000, 0)],
                      sheet="장비", extra_sheets=("메모 시트",))
    wb = load_workbook(path)
    wb.move_sheet("메모 시트", offset=-1)  # 장비 시트가 첫 시트가 아니어도 이름으로 찾는다
    wb.save(path)
    preview = imports.preview(ImportKind.ASSETS, path)
    assert preview.sheet_name == "장비" and preview.rows[0].is_ok(True)


def test_first_sheet_used_when_named_differently(imports, folder):
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [("S-1", "장비", "분류", None, None, None, 1, 0)],
                      sheet="2026년 재고")
    assert imports.preview(ImportKind.ASSETS, path).sheet_name == "2026년 재고"


def test_backup_created_before_import(imports, svc, conn, folder, tmp_path):
    svc.add_asset("CAM-1")
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [("N-1", "새 장비", "카메라", None, None, None, 1, 0)])
    backups = tmp_path / "backups" / "before_import"
    result = imports.run_import(imports.preview(ImportKind.ASSETS, path), create_categories=True,
                                backup_dir=backups)
    assert result.backup_path.parent == backups and result.backup_path.name.startswith(BACKUP_PREFIX)
    validated = inspect_backup(result.backup_path)
    assert validated.manifest.summary.counts["assets"] == 1  # 가져오기 전 상태
    validated.cleanup()


def test_backup_failure_stops_import(imports, svc, folder, tmp_path):
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [("N-1", "새 장비", "카메라", None, None, None, 1, 0)])
    blocker = tmp_path / "파일"
    blocker.write_text("x")  # 폴더 자리에 파일이 있어 백업 불가
    with pytest.raises(DomainError, match="백업을 만들지 못해"):
        imports.run_import(imports.preview(ImportKind.ASSETS, path), create_categories=True,
                           backup_dir=blocker / "sub")
    assert svc.assets.search_assets() == []


def test_changed_db_after_preview_is_refused(imports, svc, folder):
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [
        ("N-1", "새 장비", "카메라", None, None, None, 1, 0),
        ("N-2", "새 장비2", "카메라", None, None, None, 1, 0),
    ])
    preview = imports.preview(ImportKind.ASSETS, path)
    svc.add_asset("n-2")  # 미리보기 뒤 다른 곳에서 같은 관리번호 등록
    with pytest.raises(InvalidStateError, match="다시 선택"):
        imports.run_import(preview, create_categories=True, backup_dir=None)
    assert [a.asset_code for a in svc.assets.search_assets()] == ["n-2"]  # 일부만 들어가지 않음


def test_nothing_to_import(imports, folder):
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [("A", None, None, None, None, None, None, None)])
    preview = imports.preview(ImportKind.ASSETS, path)
    with pytest.raises(ValidationError, match="등록할 수 있는 행이 없습니다"):
        imports.run_import(preview, create_categories=True, backup_dir=None)


def test_blank_rows_skipped_and_row_numbers_match_excel(imports, folder):
    path = write_xlsx(folder / "a.xlsx", ASSET_HEADERS, [
        ("R-1", "a", "c", None, None, None, 1, 0),
        (None,) * 8,
        ("R-2", "b", "c", None, None, None, 1, 0),
    ])
    assert [r.row_no for r in imports.preview(ImportKind.ASSETS, path).rows] == [2, 4]


def test_1000_rows_performance(imports, svc, folder):
    rows = [(f"P-{i:04d}", f"장비 {i}", f"분류 {i % 7}", "브랜드", f"M{i}", f"S{i}", 10000 + i, 50000)
            for i in range(1000)]
    path = write_xlsx(folder / "대량.xlsx", ASSET_HEADERS, rows)
    started = time.perf_counter()
    preview = imports.preview(ImportKind.ASSETS, path)
    result = imports.run_import(preview, create_categories=True, backup_dir=None)
    elapsed = time.perf_counter() - started
    assert result.created == 1000 and len(result.created_categories) == 7
    assert len(svc.assets.search_assets()) == 1000
    assert elapsed < 15, f"1,000행 가져오기 {elapsed:.1f}초"


# --------------------------------------------------------------------------- 고객


def test_customers_import(imports, svc, folder):
    path = write_xlsx(folder / "고객.xlsx", CUSTOMER_HEADERS, [
        ("홍길동", "개인", 1012345678, "hong@example.com", "단골"),
        ("(주)촬영", "사업자", "02-123-4567", None, None),
        ("김철수", None, None, None, None),
        (None, "개인", "010", None, None),
        ("이영희", "단체", None, None, None),
    ], sheet="고객")
    preview = imports.preview(ImportKind.CUSTOMERS, path)
    ok = [r.is_ok(False) for r in preview.rows]
    assert ok == [True, True, True, False, False]
    assert "앞에 0을 붙였습니다" in preview.rows[0].notes[0]
    result = imports.run_import(preview, create_categories=False, backup_dir=None)
    assert (result.created, result.skipped) == (3, 2)
    by_name = {c.name: c for c in svc.customers.search_customers()}
    assert by_name["홍길동"].phone == "01012345678"
    assert by_name["(주)촬영"].customer_type == CustomerType.BUSINESS
    assert by_name["김철수"].customer_type == CustomerType.INDIVIDUAL


def test_customer_duplicates_by_name_and_phone(imports, svc, folder):
    svc.customers.create_customer(CustomerInput(name="홍길동", phone="010-1234-5678"))
    path = write_xlsx(folder / "고객.xlsx", CUSTOMER_HEADERS, [
        ("홍길동", None, "01012345678", None, None),     # DB와 같음(하이픈 무시)
        ("홍길동", None, "010-9999-0000", None, None),   # 전화 다르면 동명이인
        ("김철수", None, "010-1111-2222", None, None),
        ("김철수", None, "010 1111 2222", None, None),   # 파일 안 중복
        ("박영수", None, None, None, None),
        ("박영수", None, None, None, None),              # 전화 없으면 중복으로 보지 않음
    ], sheet="고객")
    rows = imports.preview(ImportKind.CUSTOMERS, path).rows
    assert [r.is_ok(False) for r in rows] == [False, True, True, False, True, True]
    assert "이미 등록" in rows[0].errors[0] and "4행과 같습니다" in rows[3].errors[0]


def test_customer_sensitive_columns_not_read(imports, svc, folder):
    path = write_xlsx(folder / "고객.xlsx", ["이름", "주민등록번호", "카드번호", "전화번호"], [
        ("홍길동", "900101-1234567", "1234-5678-9012-3456", "010-1"),
    ], sheet="고객")
    preview = imports.preview(ImportKind.CUSTOMERS, path)
    row = preview.rows[0]
    assert set(row.cells) == {"name", "phone"}
    assert any("주민등록번호" in n and "카드번호" in n for n in preview.notices)
    imports.run_import(preview, create_categories=False, backup_dir=None)
    stored = svc.customers.search_customers()[0]
    assert "900101" not in repr(stored) and "1234-5678" not in repr(stored)


def test_customer_audit_has_no_contact(imports, conn, folder):
    path = write_xlsx(folder / "고객.xlsx", CUSTOMER_HEADERS, [("홍길동", None, "010-1234-5678", "h@x.kr", None)],
                      sheet="고객")
    imports.run_import(imports.preview(ImportKind.CUSTOMERS, path), create_categories=False, backup_dir=None)
    descriptions = [r[0] for r in conn.execute("SELECT description FROM audit_events WHERE entity_type='customer'")]
    assert descriptions == ["엑셀 가져오기"]
